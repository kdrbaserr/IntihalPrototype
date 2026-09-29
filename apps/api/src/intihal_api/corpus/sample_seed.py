from __future__ import annotations

import asyncio
from dataclasses import dataclass
from datetime import UTC, datetime
from io import BytesIO

import pymupdf
from docx import Document as DocxDocument
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from starlette.concurrency import run_in_threadpool

from intihal_api.core.config import get_settings
from intihal_api.corpus.ingestion import SourceDocumentIngestionService, SourceMetadata
from intihal_api.db.models import LicenseStatus, SourceDocument
from intihal_api.db.session import AsyncSessionFactory, engine
from intihal_api.storage import ObjectStorageService, create_object_storage_service

SYNTHETIC_LICENSE = "CC0-1.0"
SYNTHETIC_RIGHTS_HOLDER = "İntihal Prototype sentetik veri üreticisi"
SYNTHETIC_CORPUS_VERSION = "v1"


@dataclass(frozen=True, slots=True)
class SampleSource:
    filename: str
    content_type: str
    content: bytes
    metadata: SourceMetadata


@dataclass(frozen=True, slots=True)
class SampleSeedResult:
    created: int
    skipped: int


def build_sample_sources() -> tuple[SampleSource, ...]:
    """Build small, original documents that exercise every supported format."""

    return (
        SampleSource(
            filename="akademik-kaynak-gosterme.txt",
            content_type="text/plain",
            content=(
                "Akademik bir metinde kullanılan düşüncenin kaynağı açıkça belirtilmelidir. "
                "Doğru atıf, okuyucunun bilginin kökenini izlemesini sağlar.\n\n"
                "Kaynakça kaydı; yazar, eser adı ve yayın bilgisini tutarlı biçimde sunar. "
                "Doğrudan alıntılar özgün ifadeyi korur ve uygun konum bilgisiyle gösterilir."
            ).encode(),
            metadata=_metadata(
                slug="akademik-kaynak-gosterme",
                title="Akademik Yazımda Kaynak Gösterme",
            ),
        ),
        SampleSource(
            filename="veri-butunlugu.pdf",
            content_type="application/pdf",
            content=_build_pdf(
                "A checksum records the identity of file content. "
                "The same bytes produce the same digest.",
                "Audit records make processing steps traceable. "
                "They do not replace access control.",
            ),
            metadata=_metadata(
                slug="veri-butunlugu",
                title="Veri Bütünlüğü ve Denetim İzleri",
            ),
        ),
        SampleSource(
            filename="tekrarlanabilir-arastirma.docx",
            content_type=(
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            ),
            content=_build_docx(),
            metadata=_metadata(
                slug="tekrarlanabilir-arastirma",
                title="Araştırmada Tekrarlanabilirlik",
            ),
        ),
    )


async def seed_sample_corpus(
    *,
    session: AsyncSession,
    storage: ObjectStorageService,
) -> SampleSeedResult:
    """Insert missing sample sources without reactivating or duplicating existing ones."""

    created = 0
    skipped = 0
    service = SourceDocumentIngestionService(storage)

    for sample in build_sample_sources():
        existing_id = await session.scalar(
            select(SourceDocument.id).where(
                SourceDocument.license_evidence_reference
                == sample.metadata.license_evidence_reference
            )
        )
        if existing_id is not None:
            skipped += 1
            continue

        source = await service.create_source(
            metadata=sample.metadata,
            filename=sample.filename,
            content_type=sample.content_type,
            stream=BytesIO(sample.content),
            session=session,
        )
        source.license_status = LicenseStatus.APPROVED
        source.license_verified_at = datetime.now(UTC)
        await session.commit()
        created += 1

    return SampleSeedResult(created=created, skipped=skipped)


def _metadata(*, slug: str, title: str) -> SourceMetadata:
    reference = f"SYNTHETIC-CORPUS-{SYNTHETIC_CORPUS_VERSION}/{slug}"
    return SourceMetadata(
        title=title,
        author="İntihal Prototype",
        publisher="Yerel geliştirme örnek havuzu",
        source_url=f"urn:intihal-prototype:synthetic:{SYNTHETIC_CORPUS_VERSION}:{slug}",
        license_name=SYNTHETIC_LICENSE,
        rights_holder=SYNTHETIC_RIGHTS_HOLDER,
        attribution_text=f"{title} — sentetik geliştirme verisi.",
        license_evidence_reference=reference,
    )


def _build_pdf(*pages: str) -> bytes:
    document = pymupdf.open()
    try:
        for content in pages:
            page = document.new_page()
            page.insert_textbox(page.rect + (50, 50, -50, -50), content, fontsize=11)
        return document.tobytes()
    finally:
        document.close()


def _build_docx() -> bytes:
    document = DocxDocument()
    document.add_heading("Tekrarlanabilir Araştırma", level=1)
    document.add_paragraph(
        "Tekrarlanabilir bir çalışma, kullanılan veri ve işlem adımlarını açıkça tanımlar."
    )
    document.add_paragraph(
        "Sürüm bilgisi ve sabit test girdileri, sonuçların daha sonra doğrulanmasını kolaylaştırır."
    )
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Kayıt"
    table.rows[0].cells[1].text = "Amaç"
    table.rows[1].cells[0].text = "Checksum"
    table.rows[1].cells[1].text = "İçerik bütünlüğü"
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


async def main() -> None:
    settings = get_settings()
    if settings.environment not in {"local", "test"}:
        raise RuntimeError("Sentetik örnek havuzu yalnızca local veya test ortamında kurulabilir.")

    storage = create_object_storage_service(settings)
    await run_in_threadpool(storage.ensure_bucket)
    try:
        async with AsyncSessionFactory() as session:
            result = await seed_sample_corpus(session=session, storage=storage)
        print(f"Synthetic corpus ready: created={result.created}, skipped={result.skipped}")
    finally:
        await engine.dispose()


if __name__ == "__main__":
    asyncio.run(main())
