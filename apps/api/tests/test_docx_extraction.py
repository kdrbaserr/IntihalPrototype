from io import BytesIO

import pytest
from docx import Document

from intihal_api.extraction import (
    InvalidDocxError,
    NoExtractableDocxTextError,
    extract_docx_text,
)


def make_docx() -> BytesIO:
    document = Document()
    document.add_heading("Thesis   title", level=1)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Source"
    table.rows[0].cells[1].text = "Score"
    document.add_paragraph("Closing paragraph")
    stream = BytesIO()
    document.save(stream)
    stream.seek(0)
    return stream


def test_extracts_paragraphs_and_tables_in_document_order() -> None:
    stream = make_docx()

    extracted = extract_docx_text(stream)

    assert extracted.text == "Thesis title\nSource Score\nClosing paragraph"
    assert extracted.page_number is None
    assert stream.tell() == 0


@pytest.mark.parametrize("content", [None, "\u200b"])
def test_rejects_docx_without_text_after_cleaning(content: str | None) -> None:
    stream = BytesIO()
    document = Document()
    if content is not None:
        document.add_paragraph(content)
    document.save(stream)

    with pytest.raises(NoExtractableDocxTextError) as captured_error:
        extract_docx_text(stream)

    assert captured_error.value.code == "no_extractable_text"
    assert stream.tell() == 0


@pytest.mark.parametrize("content", [b"", b"not a docx", b"PK\x03\x04broken"])
def test_rejects_invalid_docx(content: bytes) -> None:
    with pytest.raises(InvalidDocxError) as captured_error:
        extract_docx_text(BytesIO(content))

    assert captured_error.value.code == "invalid_docx"
