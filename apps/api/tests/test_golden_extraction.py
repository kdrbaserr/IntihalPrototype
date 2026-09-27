import json
from dataclasses import asdict
from io import BytesIO
from pathlib import Path
from typing import Any

import pymupdf
from docx import Document

from intihal_api.extraction import (
    create_text_chunks,
    extract_docx_text,
    extract_pdf_pages,
    extract_text_file,
)

GOLDEN_FIXTURE_DIRECTORY = Path(__file__).parent / "fixtures" / "golden"


def _make_pdf_fixture() -> BytesIO:
    document = pymupdf.open()
    try:
        first_page = document.new_page()
        first_page.insert_text((72, 72), "First   page. Second sentence!")
        second_page = document.new_page()
        second_page.insert_text((72, 72), "Third page?")
        return BytesIO(document.tobytes())
    finally:
        document.close()


def _make_docx_fixture() -> BytesIO:
    document = Document()
    document.add_heading("Türkçe   Başlık", level=1)
    document.add_paragraph("İlk cümle. İkinci cümle!")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Kaynak"
    table.rows[0].cells[1].text = "Puan"
    stream = BytesIO()
    document.save(stream)
    stream.seek(0)
    return stream


def _make_txt_fixture() -> BytesIO:
    text = "  Türkçe\x00\t  içerik korunur.  \r\nİkinci   cümle!  "
    return BytesIO(text.encode("utf-8"))


def _pdf_result() -> dict[str, Any]:
    extracted = extract_pdf_pages(_make_pdf_fixture())
    chunked = create_text_chunks(extracted)
    return {
        "extracted": [asdict(page) for page in extracted],
        "chunked": asdict(chunked),
    }


def _docx_result() -> dict[str, Any]:
    extracted = extract_docx_text(_make_docx_fixture())
    chunked = create_text_chunks((extracted,))
    return {"extracted": asdict(extracted), "chunked": asdict(chunked)}


def _txt_result() -> dict[str, Any]:
    extracted = extract_text_file(_make_txt_fixture())
    chunked = create_text_chunks((extracted,))
    return {"extracted": asdict(extracted), "chunked": asdict(chunked)}


def _load_golden_fixture(filename: str) -> dict[str, Any]:
    fixture_path = GOLDEN_FIXTURE_DIRECTORY / filename
    return json.loads(fixture_path.read_text(encoding="utf-8"))


def _json_result(result: dict[str, Any]) -> dict[str, Any]:
    """Compare the public JSON shape, where tuples are represented as arrays."""

    return json.loads(json.dumps(result, ensure_ascii=False))


def test_pdf_extraction_matches_golden_fixture() -> None:
    assert _json_result(_pdf_result()) == _load_golden_fixture("pdf.json")


def test_docx_extraction_matches_golden_fixture() -> None:
    assert _json_result(_docx_result()) == _load_golden_fixture("docx.json")


def test_txt_extraction_matches_golden_fixture() -> None:
    assert _json_result(_txt_result()) == _load_golden_fixture("txt.json")
